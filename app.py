import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL = "gemini-3.8-flash"
# Tried in order if the main model is overloaded (503) or rate-limited (429).
FALLBACK_MODELS = ["gemini-flash-latest"]
RETRIES_PER_MODEL = 3

st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")


# ---------- Text extraction ----------
def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        doc = Document(uploaded_file)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    return uploaded_file.read().decode("utf-8", errors="ignore")


# ---------- Rule-based checks (no AI needed) ----------
ACTION_VERBS = {
    "led", "built", "developed", "designed", "created", "managed", "improved",
    "increased", "reduced", "launched", "implemented", "optimized", "delivered",
    "analyzed", "automated", "mentored", "negotiated", "achieved", "streamlined",
}
SECTIONS = {
    "Experience": r"\b(experience|employment|work history)\b",
    "Education": r"\beducation\b",
    "Skills": r"\bskills\b",
    "Summary/Objective": r"\b(summary|objective|profile)\b",
    "Projects/Certifications": r"\b(projects|certifications?)\b",
}


def rule_based_score(text: str):
    lower = text.lower()
    words = re.findall(r"\w+", text)
    checks = []  # (label, points earned, points possible, tip)

    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text))
    has_phone = bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text))
    pts = (10 if has_email else 0) + (10 if has_phone else 0)
    checks.append(("Contact info", pts, 20, "Include a clear email and phone number."))

    found = [s for s, pat in SECTIONS.items() if re.search(pat, lower)]
    pts = round(30 * len(found) / len(SECTIONS))
    missing = [s for s in SECTIONS if s not in found]
    tip = "Add standard headings: " + ", ".join(missing) + "." if missing else "All standard sections found."
    checks.append(("Standard sections", pts, 30, tip))

    n = len(words)
    if 300 <= n <= 1000:
        pts, tip = 15, "Length looks good."
    elif 150 <= n < 300 or 1000 < n <= 1400:
        pts, tip = 8, "Aim for roughly 300-1000 words (1-2 pages)."
    else:
        pts, tip = 2, "Resume is too short or too long; aim for 300-1000 words."
    checks.append(("Length", pts, 15, tip))

    verbs = sum(1 for w in words if w.lower() in ACTION_VERBS)
    pts = min(15, verbs * 2)
    checks.append(("Action verbs", pts, 15, "Start bullets with strong verbs (led, built, improved...)."))

    metrics = len(re.findall(r"\d+\s?%|\$\s?\d+|\b\d{2,}\b", text))
    pts = min(20, metrics * 3)
    checks.append(("Quantified results", pts, 20, "Add numbers: percentages, revenue, users, time saved."))

    total = sum(c[1] for c in checks)
    return total, checks


# ---------- AI analysis ----------
PROMPT = """You are an expert ATS (Applicant Tracking System) and resume reviewer.
Evaluate the resume below{jd_clause}. Respond ONLY with JSON in this exact shape:
{{
  "ai_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["..."],
  "weaknesses": ["..."],
  "suggestions": ["<specific, actionable improvement>"],
  "missing_keywords": ["..."],
  "job_match_percent": <integer 0-100 or null if no job description given>
}}
Be honest and specific; do not inflate the score.

RESUME:
\"\"\"
{resume}
\"\"\"
{jd_block}"""


def analyze_with_ai(api_key: str, resume: str, job_desc: str) -> dict:
    client = genai.Client(api_key=api_key)
    jd_clause = " against the job description" if job_desc.strip() else " for general ATS-friendliness"
    jd_block = f'\nJOB DESCRIPTION:\n"""\n{job_desc}\n"""' if job_desc.strip() else ""
    prompt = PROMPT.format(jd_clause=jd_clause, resume=resume[:15000], jd_block=jd_block)
    config = types.GenerateContentConfig(
        response_mime_type="application/json", temperature=0.2
    )
    resp = None
    last_err = None
    for model in [MODEL] + FALLBACK_MODELS:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                resp = client.models.generate_content(
                    model=model, contents=prompt, config=config
                )
                break
            except Exception as e:
                last_err = e
                msg = str(e)
                temporary = any(k in msg for k in ("503", "429", "UNAVAILABLE", "RESOURCE_EXHAUSTED", "overloaded"))
                if not temporary:
                    raise  # real error (bad key, 404 model, etc.): don't retry
                time.sleep(2 ** attempt)  # wait 1s, 2s, 4s
        if resp is not None:
            break
    if resp is None:
        raise last_err
    raw = (resp.text or "").strip()
    raw = re.sub(r"^```(?:json)?|```$", "", raw, flags=re.MULTILINE).strip()
    return json.loads(raw)


def get_api_key() -> str:
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "")


# ---------- UI ----------
st.title("📄 ATS Resume Checker")
st.caption("Upload your resume and get an ATS score with suggestions to improve it.")

api_key = get_api_key()
if not api_key:
    api_key = st.text_input("Gemini API key", type="password")

uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
job_desc = st.text_area("Job description (optional, for a match score)", height=150)

if st.button("Check ATS score", type="primary", disabled=uploaded is None):
    if not api_key:
        st.error("Please provide a Gemini API key.")
        st.stop()

    with st.spinner("Reading your resume..."):
        text = extract_text(uploaded).strip()

    if len(text) < 100:
        st.error("Couldn't extract enough text. If your PDF is a scanned image, ATS systems can't read it either — export a text-based PDF or use DOCX.")
        st.stop()

    rule_score, checks = rule_based_score(text)

    try:
        with st.spinner("Analyzing with AI..."):
            ai = analyze_with_ai(api_key, text, job_desc)
    except Exception as e:
        if any(k in str(e) for k in ("503", "429", "UNAVAILABLE")):
            st.error("Google's AI servers are busy right now. Please wait a minute and try again.")
        else:
            st.error(f"AI analysis failed: {e}")
        st.stop()

    ai_score = max(0, min(100, int(ai.get("ai_score", 0))))
    final = round(0.4 * rule_score + 0.6 * ai_score)

    st.divider()
    st.metric("Overall ATS score", f"{final} / 100")
    st.progress(final / 100)
    c1, c2 = st.columns(2)
    c1.metric("Format & structure", f"{rule_score} / 100")
    c2.metric("AI content review", f"{ai_score} / 100")
    if ai.get("job_match_percent") is not None:
        st.metric("Job match", f"{ai['job_match_percent']}%")

    st.subheader("Summary")
    st.write(ai.get("summary", ""))

    st.subheader("Format checks")
    for label, pts, possible, tip in checks:
        icon = "✅" if pts >= possible * 0.8 else "⚠️"
        st.write(f"{icon} **{label}** — {pts}/{possible}. {tip if pts < possible * 0.8 else ''}")

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in ai.get("strengths", []):
            st.write(f"- {s}")
    with right:
        st.subheader("Weaknesses")
        for w in ai.get("weaknesses", []):
            st.write(f"- {w}")

    st.subheader("Suggestions")
    for s in ai.get("suggestions", []):
        st.write(f"- {s}")

    if ai.get("missing_keywords"):
        st.subheader("Missing keywords")
        st.write(", ".join(ai["missing_keywords"]))
